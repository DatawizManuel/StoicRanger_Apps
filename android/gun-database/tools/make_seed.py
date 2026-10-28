#!/usr/bin/env python3
"""Build app data files.
  make_seed.py demo <examples dir> <out.json>          demo database from 05_tools/examples
  make_seed.py xlsx <database.xlsx> <out.json>         backup file (app format) from the PC registry database.xlsx
"""
import glob, json, os, sys
from datetime import datetime

KEYS = ['no','date','time','shooter','purpose','discipline','position','series','rounds','gun','barrel','optic','caliber','ammo','bullet',
        'v0','v0sd','zero','dialE','dialW','range','coords','elev','dist','setting','temp','hum','press','wind','light','preset','bd','d10','w','remarks']

def base():
    now = datetime.now().isoformat(timespec='seconds')
    return {"schema": 1, "created": now, "updated": now, "weapons": [], "barrels": [], "munitions": [], "shooters": [], "roundsLog": [], "sheets": [],
            "settings": {"preset": "issf300", "claudeModel": "claude-opus-5-5", "theme": "auto", "sheetPrefix": "ST"}}

def demo(src, out):
    db = base()
    db["weapons"].append({"id": "W-001", "name": "Tikka T3x TAC A1", "aliases": "Tikka", "make": "Tikka", "model": "T3x TAC A1", "caliber": "6.5 Creedmoor",
                          "barrelLengthMm": 610, "optic": "Steiner M7Xi 4-28x56 MSR-2", "status": "Active", "roundsBeforeLog": 0, "notes": "Example rifle"})
    bars = json.load(open(os.path.join(src, "_barrels.json")))
    for i, b in enumerate(bars):
        db["barrels"].append({"id": f"B-{i+1:03d}", "weaponId": "W-001", "name": b["barrel"], "installed": b["installed"], "removed": b["removed"] or "",
                              "roundsAtInstall": b["rounds_at_install"], "notes": "Example"})
    bid = {b["name"]: b["id"] for b in db["barrels"]}
    for i, l in enumerate(json.load(open(os.path.join(src, "_loads.json")))):
        db["munitions"].append({"id": f"M-{i+1:03d}", "name": l["ammo"], "type": "Handload" if l["type"] == "handload" else "Factory", "caliber": "6.5 Creedmoor",
                                "bullet": l.get("bullet", ""), "bulletWeightG": l.get("bullet_g"), "powder": l.get("powder", ""), "chargeG": l.get("charge_g"),
                                "primer": l.get("primer", ""), "cbtoMm": l.get("cbto_mm"), "status": "Active"})
    mid = {m["name"]: m["id"] for m in db["munitions"]}
    db["shooters"].append({"id": "S-001", "name": "Example shooter", "status": "Active"})
    for i, r in enumerate(json.load(open(os.path.join(src, "_rounds-log.json")))):
        db["roundsLog"].append({"id": f"R-{i+1:03d}", "date": r["date"], "weaponId": "W-001", "barrelId": bid.get(r["barrel"]), "rounds": r["rounds"], "note": r["note"]})
    for f in sorted(glob.glob(os.path.join(src, "2026*.json"))):
        r = json.load(open(f))
        s = {k: str(r.get(k, "") if r.get(k) is not None else "") for k in KEYS}
        s["time"] = s["time"].zfill(5) if s["time"] else ""
        s.update(id="SH-" + r["no"], hits=r["hits"], weaponId="W-001", barrelId=bid.get(r["barrel"]), munitionId=mid.get(r["ammo"]), shooterId="S-001", shooter="Example shooter")
        db["sheets"].append(s)
    json.dump(db, open(out, "w"), ensure_ascii=False, separators=(",", ":"))

def xlsx(path, out):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    def rows(name):
        if name not in wb.sheetnames: return []
        it = wb[name].iter_rows(values_only=True)
        head = [str(h).strip() if h else "" for h in next(it)]
        return [dict(zip(head, r)) for r in it if any(v not in (None, "") for v in r)]
    def d(v): return v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else (str(v) if v not in (None, "") else "")
    def t(v): return "" if v is None else str(v)
    db = base()
    for r in rows("Weapons"):
        db["weapons"].append({"id": r["ID"], "name": t(r.get("Name")), "aliases": t(r.get("Aliases (as written on sheets)")), "make": t(r.get("Make")), "model": t(r.get("Model")),
            "caliber": t(r.get("Caliber")), "barrelLengthMm": r.get("Barrel length mm"), "twist": t(r.get("Twist")), "thread": t(r.get("Muzzle thread")), "optic": t(r.get("Optic")),
            "mount": t(r.get("Mount")), "accessories": t(r.get("Accessories")), "weightKg": r.get("Weight bare kg"), "serial": t(r.get("Serial no.")), "added": d(r.get("Added")),
            "status": t(r.get("Status")) or "Active", "notes": t(r.get("Notes")), "roundsBeforeLog": r.get("Rounds before log") or 0, "countedOn": d(r.get("Counted on"))})
    for r in rows("Barrels"):
        db["barrels"].append({"id": r["ID"], "weaponId": r.get("Weapon ID"), "name": t(r.get("Barrel")), "installed": d(r.get("Installed")), "removed": d(r.get("Removed")),
            "roundsAtInstall": r.get("Rounds at install") or 0, "notes": t(r.get("Notes"))})
    for r in rows("Munitions"):
        db["munitions"].append({"id": r["ID"], "name": t(r.get("Name")), "aliases": t(r.get("Aliases (as written on sheets)")), "type": t(r.get("Type")) or "Factory",
            "caliber": t(r.get("Caliber")), "manufacturer": t(r.get("Manufacturer")), "bullet": t(r.get("Bullet")), "bulletWeightG": r.get("Bullet weight g"), "bcG7": r.get("BC G7"),
            "nominalV0": r.get("Nominal V0 m/s"), "powder": t(r.get("Powder")), "chargeG": r.get("Charge g"), "primer": t(r.get("Primer")), "brass": t(r.get("Brass")),
            "cbtoMm": r.get("CBTO mm"), "coalMm": r.get("COAL mm"), "lot": t(r.get("Lot")), "added": d(r.get("Added")), "status": t(r.get("Status")) or "Active", "notes": t(r.get("Notes"))})
    for r in rows("Shooters"):
        db["shooters"].append({"id": r["ID"], "name": t(r.get("Name")), "aliases": t(r.get("Aliases (as written on sheets)")), "added": d(r.get("Added")),
            "status": t(r.get("Status")) or "Active", "notes": t(r.get("Notes"))})
    for i, r in enumerate(rows("Rounds_Log")):
        db["roundsLog"].append({"id": f"R-{i+1:03d}", "date": d(r.get("Date")), "weaponId": r.get("Weapon ID"), "barrelId": r.get("Barrel ID"), "rounds": r.get("Rounds") or 0, "note": t(r.get("Note"))})
    db["app"] = "Range Logbook backup"
    json.dump(db, open(out, "w"), ensure_ascii=False, indent=1)

if __name__ == "__main__":
    {"demo": demo, "xlsx": xlsx}[sys.argv[1]](sys.argv[2], sys.argv[3])
